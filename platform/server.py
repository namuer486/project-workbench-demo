"""Spreadsheet-first project workbench. Run: python server.py."""
from __future__ import annotations

import base64
import csv
import hashlib
import hmac
import io
import ipaddress
import json
import os
import secrets
import sqlite3
import threading
import time
import zipfile
from contextlib import contextmanager
from datetime import date, datetime, timedelta
from http.cookies import SimpleCookie
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

from ipd import contract, report_snapshot

ROOT = Path(__file__).resolve().parent
DEFAULT_TAGS = ["新鲜感", "目标感", "成长感", "乐趣性", "社交感", "时间成本"]
FIELDS = {"id": "问题编号", "title": "问题内容", "owner": "负责人", "status": "完成状态",
          "date": "提出日期", "solution": "解决方案", "priority": "优先级", "category": "功能分类",
          "tags": "衡量标签", "version": "版本", "resolvedAt": "解决日期"}
ALIASES = {"id": ["编号", "问题编号", "ID", "序号"], "title": ["问题内容", "问题描述", "标题", "问题", "title"],
           "owner": ["负责人", "责任人", "owner"], "status": ["完成状态", "状态", "解决状态", "status"],
           "date": ["提出日期", "提出时间", "日期", "创建时间", "date"], "solution": ["解决方案", "处理方案", "方案", "solution"],
           "priority": ["优先级", "priority"], "category": ["功能分类", "分类", "模块", "category"],
           "tags": ["衡量标签", "标签", "维度", "tags"], "version": ["版本", "所属版本", "version"],
           "resolvedAt": ["解决日期", "完成日期", "解决时间", "resolvedAt"]}
STATUS = {"未解决": "未解决", "待处理": "未解决", "未开始": "未解决", "未完成": "未解决", "open": "未解决",
          "进行中": "进行中", "处理中": "进行中", "修复中": "进行中", "in progress": "进行中",
          "待验证": "待验证", "待验收": "待验证", "已解决": "已解决", "已完成": "已解决", "完成": "已解决",
          "已关闭": "已解决", "done": "已解决", "closed": "已解决"}
MAX_FILE = 10 * 1024 * 1024
MAX_ROWS = 20000


class Error(Exception):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def dumps(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def now():
    return datetime.now().astimezone().isoformat(timespec="seconds")


def digest(value):
    return hashlib.sha256(value.encode()).hexdigest()


def text(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip() if value is not None else ""


def parse_date(value):
    value = text(value)
    if not value:
        return ""
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("日期需为 YYYY-MM-DD（Excel 日期单元格也可）")


def parse_table(filename, raw, sheet=None):
    if len(raw) > MAX_FILE:
        raise Error("文件不能超过 10 MB")
    suffix = Path(filename).suffix.lower()
    sheets = []
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 80 * 1024 * 1024:
                    raise Error("Excel 解压体积过大，请精简后导入")
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                sheets = book.sheetnames
                if sheet and sheet not in sheets:
                    raise Error("工作表不存在")
                selected = sheet or sheets[0]
                ws = book[selected]
                if ws.max_column and ws.max_column > 100:
                    raise Error("最多支持 100 列，请删除多余格式列")
                rows = []
                for row in ws.iter_rows():
                    if any(cell.data_type == "f" for cell in row):
                        raise Error("表中含公式，请复制并粘贴为值后导入，避免过期的公式计算结果")
                    rows.append([text(cell.value) for cell in row])
                    if len(rows) > MAX_ROWS + 1:
                        raise Error("每次最多导入 20,000 行")
            finally:
                book.close()
        except Error:
            raise
        except Exception as exc:
            raise Error("无法读取 Excel，请使用未加密的 .xlsx 文件") from exc
    elif suffix == ".csv":
        selected = "CSV"
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                content = raw.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise Error("CSV 请使用 UTF-8 或 GB18030 编码") from exc
        rows = []
        try:
            for row in csv.reader(io.StringIO(content), strict=True):
                rows.append([text(v) for v in row])
                if len(rows) > MAX_ROWS + 1:
                    raise Error("每次最多导入 20,000 行")
        except csv.Error as exc:
            raise Error("CSV 格式不正确") from exc
    else:
        raise Error("支持 .xlsx 和 .csv；旧版 .xls 请另存为 .xlsx")
    if not rows or not any(rows[0]):
        raise Error("第一行必须是列名，文件不能为空")
    headers = rows[0]
    if len(headers) > 100 or any(not h for h in headers) or len(set(headers)) != len(headers):
        raise Error("第一行列名必须非空且不重复，最多 100 列")
    records = []
    for number, values in enumerate(rows[1:], 2):
        if not any(values):
            continue
        if len(values) > len(headers):
            raise Error(f"第 {number} 行超出表头列数")
        if any(len(v) > 20000 for v in values):
            raise Error(f"第 {number} 行单元格超过 20,000 字")
        records.append((number, dict(zip(headers, values + [""] * (len(headers) - len(values))))))
    if not records:
        raise Error("文件没有数据行")
    return headers, records, sheets, selected


def validate_config(payload):
    name = text(payload.get("name"))
    if not name or len(name) > 80:
        raise Error("项目名称必填，最多 80 字")
    try:
        start = parse_date(payload.get("startDate"))
        if not start:
            raise ValueError()
    except ValueError as exc:
        raise Error("请填写有效的项目起始日期") from exc
    tags = payload.get("tags", DEFAULT_TAGS)
    if not isinstance(tags, list) or not 1 <= len(tags) <= 20 or any(not isinstance(t, str) or not t.strip() or len(t) > 30 for t in tags):
        raise Error("衡量标签需为 1–20 项，每项最多 30 字")
    tags = [t.strip() for t in tags]
    if len(tags) != len(set(tags)):
        raise Error("衡量标签不能重复")
    versions = payload.get("versions", [])
    if not isinstance(versions, list) or len(versions) > 50:
        raise Error("版本最多 50 个")
    clean_versions = []
    try:
        for v in versions:
            n, d = text(v["name"]), parse_date(v["startDate"])
            if not n or not d or len(n) > 40:
                raise ValueError()
            clean_versions.append({"name": n, "startDate": d})
    except (TypeError, KeyError, ValueError) as exc:
        raise Error("每个版本需要名称和有效起始日期") from exc
    if len({v["name"] for v in clean_versions}) != len(clean_versions) or len({v["startDate"] for v in clean_versions}) != len(clean_versions):
        raise Error("版本名称与起始日期不能重复")
    return {"name": name, "startDate": start, "tags": tags,
            "versions": sorted(clean_versions, key=lambda v: v["startDate"]), "stage": text(payload.get("stage"))[:60]}


def normalize(records, mapping, config):
    if not isinstance(mapping, dict) or not mapping.get("id") or not mapping.get("title"):
        raise Error("必须映射「问题编号」和「问题内容」，编号用于重复识别和更新")
    if any(k not in FIELDS for k in mapping):
        raise Error("存在未知映射字段")
    selected = [v for v in mapping.values() if v]
    if len(selected) != len(set(selected)):
        raise Error("一个表格列只能映射到一个字段")
    headers = set(records[0][1])
    if any(v not in headers for v in selected):
        raise Error("映射列不存在，请重新选择")
    items, errors, seen = [], [], set()
    for number, record in records:
        item = {key: text(record.get(mapping.get(key, ""), "")) for key in FIELDS}
        row_errors = []
        if not item["id"] or not item["title"]:
            row_errors.append("问题编号和问题内容不能为空")
        if len(item["id"]) > 120:
            row_errors.append("问题编号最多 120 字")
        if item["id"] in seen:
            row_errors.append("问题编号重复：" + item["id"])
        seen.add(item["id"])
        status = item["status"].lower()
        if status and status not in STATUS:
            row_errors.append("无法识别状态：" + item["status"] + "（支持未解决/进行中/待验证/已解决）")
        item["status"] = STATUS.get(status, "未解决")
        for key in ("date", "resolvedAt"):
            try:
                item[key] = parse_date(item[key])
            except ValueError as exc:
                row_errors.append(FIELDS[key] + "：" + str(exc))
        if item["date"] and item["resolvedAt"] and item["resolvedAt"] < item["date"]:
            row_errors.append("解决日期不能早于提出日期")
        import re
        item["tags"] = list(dict.fromkeys(t.strip() for t in re.split(r"[,，;；|、]", item["tags"]) if t.strip()))
        unknown = set(item["tags"]) - set(config["tags"])
        if unknown:
            row_errors.append("未配置的衡量标签：" + "、".join(sorted(unknown)))
        item["owner"] = item["owner"] or "未分配"
        item["category"] = item["category"] or "未分类"
        item["priority"] = item["priority"] or "未标注"
        if row_errors:
            errors.append({"row": number, "messages": row_errors})
        items.append(item)
    return items, errors


class Store:
    def __init__(self, directory):
        self.directory = Path(directory)
        self.directory.mkdir(parents=True, exist_ok=True)
        self.path = self.directory / "workbench.sqlite3"
        with self.connect() as db:
            db.executescript("""
            PRAGMA journal_mode=WAL;
            CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY, config TEXT NOT NULL, revision INTEGER NOT NULL DEFAULT 0,
                mapping TEXT NOT NULL DEFAULT '{}', updated TEXT NOT NULL, share_hash TEXT);
            CREATE TABLE IF NOT EXISTS reports(project_id TEXT PRIMARY KEY REFERENCES projects(id), document TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS report_versions(project_id TEXT NOT NULL REFERENCES projects(id), revision INTEGER NOT NULL,
                created TEXT NOT NULL, document TEXT NOT NULL, PRIMARY KEY(project_id,revision));
            CREATE TABLE IF NOT EXISTS items(project_id TEXT NOT NULL REFERENCES projects(id), id TEXT NOT NULL, data TEXT NOT NULL,
                PRIMARY KEY(project_id,id));
            CREATE TABLE IF NOT EXISTS imports(id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id), filename TEXT NOT NULL,
                created TEXT NOT NULL, summary TEXT NOT NULL, before_data TEXT NOT NULL, after_data TEXT NOT NULL);
            """)

    @contextmanager
    def connect(self):
        db = sqlite3.connect(self.path, timeout=15)
        db.row_factory = sqlite3.Row
        db.execute("PRAGMA foreign_keys=ON")
        try:
            with db:
                yield db
        finally:
            db.close()

    def project(self, project_id, db=None):
        if db is None:
            with self.connect() as conn:
                return self.project(project_id, conn)
        row = db.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
        if not row:
            raise Error("项目不存在", 404)
        return {"id": row["id"], **json.loads(row["config"]), "revision": row["revision"],
                "mapping": json.loads(row["mapping"]), "updated": row["updated"], "sharing": bool(row["share_hash"])}

    def create(self, payload):
        config = validate_config(payload)
        project_id = secrets.token_hex(8)
        with self.connect() as db:
            db.execute("INSERT INTO projects(id,config,updated) VALUES(?,?,?)", (project_id, dumps(config), now()))
        return self.project(project_id)

    def update(self, project_id, payload):
        if self.project(project_id).get("mode") == "report":
            raise Error("项目拆解请通过新版 JSON 更新，不能单独修改配置")
        config = validate_config(payload)
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            project = self.project(project_id, db)
            if payload.get("revision") != project["revision"]:
                raise Error("项目已被更新，请刷新后再保存", 409)
            used = set()
            for row in db.execute("SELECT data FROM items WHERE project_id=?", (project_id,)):
                used.update(json.loads(row[0])["tags"])
            if used - set(config["tags"]):
                raise Error("已有问题使用了这些标签，请先在源表中调整并导入：" + "、".join(sorted(used - set(config["tags"]))))
            db.execute("UPDATE projects SET config=?,revision=revision+1,updated=? WHERE id=?", (dumps(config), now(), project_id))
        return self.project(project_id)

    def delete(self, project_id, payload):
        if payload.get("confirmed") is not True:
            raise Error("请先确认删除项目")
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            project = self.project(project_id, db)
            if payload.get("revision") != project["revision"]:
                raise Error("项目已被更新，请刷新后重新确认删除", 409)
            for table in ("imports", "items", "report_versions", "reports"):
                db.execute(f"DELETE FROM {table} WHERE project_id=?", (project_id,))
            db.execute("DELETE FROM projects WHERE id=?", (project_id,))
        return {"ok": True, "projectId": project_id}

    def preview_report(self, report):
        errors = contract.validate(report)
        if errors:
            raise Error("项目 JSON 校验失败：" + "；".join(errors[:20]))
        if len(dumps(report).encode("utf-8")) > MAX_FILE:
            raise Error("项目 JSON 不能超过 10 MB", 413)
        project_id = report["project"]["id"]
        with self.connect() as db:
            exists = db.execute("SELECT id FROM projects WHERE id=?", (project_id,)).fetchone()
            project = self.project(project_id, db) if exists else None
        if project and project.get("mode") != "report":
            raise Error("此项目 ID 已被表格项目使用，请换一个项目 ID", 409)
        return {"kind": "report", "projectId": project_id, "revision": project["revision"] if project else None,
                "previousName": project["name"] if project else None, "report": report}

    def save_report(self, preview):
        report, project_id = preview["report"], preview["projectId"]
        snapshot = report_snapshot(report, now())
        if report["project"]["id"] != project_id:
            raise Error("项目 ID 不匹配")
        config = {k: v for k, v in snapshot["project"].items() if k not in ("id", "revision", "updated")}
        config["mode"] = "report"
        document = dumps(report)
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            row = db.execute("SELECT revision FROM projects WHERE id=?", (project_id,)).fetchone()
            if (row["revision"] if row else None) != preview["revision"]:
                raise Error("此项目已被其他人更新，请重新导入预览", 409)
            stamp = now()
            revision = (row["revision"] if row else 0) + 1
            if row:
                if self.project(project_id, db).get("mode") != "report":
                    raise Error("不能用项目 JSON 覆盖表格项目", 409)
                db.execute("UPDATE projects SET config=?,revision=?,updated=? WHERE id=?", (dumps(config), revision, stamp, project_id))
            else:
                db.execute("INSERT INTO projects(id,config,revision,updated) VALUES(?,?,?,?)", (project_id, dumps(config), revision, stamp))
            db.execute("INSERT INTO reports VALUES(?,?) ON CONFLICT(project_id) DO UPDATE SET document=excluded.document", (project_id, document))
            db.execute("INSERT INTO report_versions VALUES(?,?,?,?)", (project_id, revision, stamp, document))
            db.execute("DELETE FROM items WHERE project_id=?", (project_id,))
            db.executemany("INSERT INTO items VALUES(?,?,?)", [(project_id, i["id"], dumps(i)) for i in snapshot["items"]])
        return {"projectId": project_id, "revision": revision}

    def edit_problem(self, project_id, payload):
        """Edit the stored document, preserving source evidence and optimistic locking."""
        snapshot = self.snapshot(project_id)
        report = snapshot.get("report")
        if not report or report["schemaVersion"] != "2.0":
            raise Error("请在已保存的新版项目拆解中维护项目问题")
        revision = snapshot["project"]["revision"]
        if payload.get("revision") != revision:
            raise Error("项目已被其他人更新。请保留当前填写内容，刷新后重新编辑。", 409)
        operation = payload.get("operation")
        if operation not in ("create", "update", "delete"):
            raise Error("不支持的问题操作")
        pid = payload.get("id")
        old = next((p for p in report["projectProblems"] if p["id"] == pid), None)
        if operation != "create" and old is None:
            raise Error("项目问题不存在", 404)
        if operation == "delete":
            if payload.get("confirmed") is not True:
                raise Error("请先确认删除问题")
            report["projectProblems"] = [p for p in report["projectProblems"] if p["id"] != pid]
            for insight in report["insights"]:
                insight["projectProblemIds"] = [i for i in insight["projectProblemIds"] if i != pid]
        else:
            fields = payload.get("fields")
            limits = {"title": 200, "owner": 120, "version": 80, "category": 100,
                      "severity": 50, "source": 200, "date": 10, "status": 10,
                      "goal": 4000, "solution": 4000, "review": 4000, "ksf": 40000}
            if not isinstance(fields, dict) or set(fields) != set(limits) | {"dimensionIds", "caseIds", "issueIds"}:
                raise Error("问题表单字段不完整，请刷新页面后重试")
            for key, limit in limits.items():
                if not isinstance(fields[key], str) or len(fields[key]) > limit:
                    raise Error("字段格式不正确或文字过长：" + key)
                fields[key] = fields[key].strip()
            for key in ("dimensionIds", "caseIds", "issueIds"):
                if not isinstance(fields[key], list) or any(not isinstance(v, str) for v in fields[key]):
                    raise Error("关联字段格式不正确")
            if not fields["title"]:
                raise Error("请填写事项标题")
            if operation == "create" and not fields["goal"]:
                raise Error("请填写目标成果")
            if fields["status"] == "已解决" and not fields["solution"]:
                raise Error("已解决问题请填写解决方案")
            pid = old["id"] if old else "manual-" + secrets.token_hex(8)

        stamp = now()
        # The author is not inferred: the local deployment has no individual accounts.
        source = next((s for s in report["sources"] if s["title"] == "工作台在线维护记录" and s["type"] == "other"), None)
        if source is None:
            source = {"id": "online-" + secrets.token_hex(8), "title": "工作台在线维护记录", "type": "other"}
            report["sources"].append(source)
        def evidence(label, text):
            ids = []
            for n in range(0, len(text), 550):
                part = text[n:n + 550]
                if not part.strip():
                    continue
                eid = "online-" + secrets.token_hex(12)
                report["evidence"].append({"id": eid, "sourceId": source["id"],
                    "location": f"{stamp} / {pid} / {label} / 第{n // 550 + 1}段", "excerpt": part})
                ids.append(eid)
            return ids
        audit = evidence("操作记录", {"create": "人工新增", "update": "人工编辑", "delete": "人工删除"}[operation] + "项目问题：" + (old["title"] if operation == "delete" else fields["title"]))
        if operation != "delete":
            item = dict(old) if old else {"id": pid, "evidenceRefs": []}
            item["origin"] = "manual"
            for key in ("title", "date", "status", "owner", "version", "category", "severity", "source"):
                item[key] = fields[key] or None
            for key in ("dimensionIds", "caseIds", "issueIds"):
                item[key] = fields[key]
            def manual_content(txt, label):
                ev = evidence(label, txt)
                audit.extend(ev)
                return {"text": txt, "origin": "manual", "evidenceRefs": ev}
            for key in ("goal", "solution", "review"):
                existing = old.get(key) if old else None
                txt = fields[key]
                item[key] = existing if existing and existing["text"] == txt else manual_content(txt, key) if txt else None
            existing_ksf = old["ksf"] if old else []
            if fields["ksf"] == "\n".join(v["text"] for v in existing_ksf):
                item["ksf"] = existing_ksf
            else:
                item["ksf"] = [manual_content(line.strip(), "KSF") for line in fields["ksf"].splitlines() if line.strip()]
            audit.extend(evidence("基本信息与关联", dumps({k: item[k] for k in ("title", "date", "status", "owner", "version", "category", "severity", "source", "dimensionIds", "caseIds", "issueIds")})))
            item["evidenceRefs"] = list(dict.fromkeys(item["evidenceRefs"] + audit))
            if old:
                report["projectProblems"] = [item if p["id"] == pid else p for p in report["projectProblems"]]
            else:
                report["projectProblems"].append(item)
        report["reviewStatus"] = "draft"
        errors = contract.validate(report)
        if errors:
            raise Error("问题未保存：" + "；".join(errors[:10]))
        if len(dumps(report).encode("utf-8")) > MAX_FILE:
            raise Error("项目数据不能超过 10 MB", 413)
        result = self.save_report({"report": report, "projectId": project_id, "revision": revision})
        return {**result, "problemId": pid}

    def snapshot(self, project_id):
        with self.connect() as db:
            db.execute("BEGIN")
            project = self.project(project_id, db)
            if project.get("mode") == "report":
                row = db.execute("SELECT document FROM reports WHERE project_id=?", (project_id,)).fetchone()
                result = report_snapshot(json.loads(row[0]), project["updated"])
                result["project"] = project
                result["history"] = [dict(h) for h in db.execute("SELECT revision,created FROM report_versions WHERE project_id=? ORDER BY revision DESC LIMIT 30", (project_id,))]
                return result
            items = [json.loads(row[0]) for row in db.execute("SELECT data FROM items WHERE project_id=? ORDER BY id", (project_id,))]
            history = [dict(row) for row in db.execute("SELECT id,filename,created,summary FROM imports WHERE project_id=? ORDER BY created DESC,rowid DESC LIMIT 30", (project_id,))]
        for entry in history:
            entry["summary"] = json.loads(entry["summary"])
        for item in items:
            item["effectiveVersion"] = item["version"] or next((v["name"] for v in reversed(project["versions"]) if item["date"] and v["startDate"] <= item["date"]), "未划分")
            if item["date"]:
                days = (date.fromisoformat(item["date"]) - date.fromisoformat(project["startDate"])).days
                item["week"] = f"第 {days // 7 + 1} 周" if days >= 0 else "开始日期之前"
                item["weekIndex"] = days // 7 if days >= 0 else -1
            else:
                item["week"], item["weekIndex"] = "未标注日期", -2
        return {"project": project, "items": items, "history": history}

    def preview(self, project_id, filename, records, mapping):
        snapshot = self.snapshot(project_id)
        project = snapshot["project"]
        if project.get("mode") == "report":
            raise Error("项目拆解请导入完整 JSON，避免明细与提炼关系不同步")
        incoming, errors = normalize(records, mapping, project)
        before = {i["id"]: {k: i[k] for k in FIELDS} for i in snapshot["items"]}
        counts = {"added": 0, "updated": 0, "unchanged": 0, "missing": 0}
        changes = []
        for item in incoming:
            old = before.get(item["id"])
            kind = "added" if old is None else "unchanged" if old == item else "updated"
            counts[kind] += 1
            if kind != "unchanged":
                changes.append({"kind": kind, "id": item["id"], "title": item["title"],
                                "fields": [{"field": FIELDS[k], "before": old[k], "after": item[k]} for k in FIELDS if old and old[k] != item[k]]})
        ids = {i["id"] for i in incoming}
        missing = [i for key, i in before.items() if key not in ids]
        counts["missing"] = len(missing)
        return {"projectId": project_id, "revision": project["revision"], "filename": filename, "mapping": mapping,
                "items": incoming, "errors": errors, "counts": counts, "changes": changes[:200],
                "missingItems": [{"id": i["id"], "title": i["title"]} for i in missing[:100]], "total": len(incoming)}

    def commit(self, preview):
        if preview["errors"]:
            raise Error("请先修复源表中的错误，再重新导入")
        project_id = preview["projectId"]
        with self.connect() as db:
            db.execute("BEGIN IMMEDIATE")
            project = self.project(project_id, db)
            if project["revision"] != preview["revision"]:
                raise Error("预览后项目数据或配置发生变化，请重新预览", 409)
            before = [json.loads(row[0]) for row in db.execute("SELECT data FROM items WHERE project_id=?", (project_id,))]
            for item in preview["items"]:
                db.execute("INSERT INTO items(project_id,id,data) VALUES(?,?,?) ON CONFLICT(project_id,id) DO UPDATE SET data=excluded.data", (project_id, item["id"], dumps(item)))
            after = [json.loads(row[0]) for row in db.execute("SELECT data FROM items WHERE project_id=?", (project_id,))]
            batch_id = secrets.token_hex(12)
            db.execute("INSERT INTO imports VALUES(?,?,?,?,?,?,?)", (batch_id, project_id, preview["filename"], now(), dumps(preview["counts"]), dumps(before), dumps(after)))
            db.execute("UPDATE projects SET revision=revision+1,mapping=?,updated=? WHERE id=?", (dumps(preview["mapping"]), now(), project_id))
        return {"batchId": batch_id, "counts": preview["counts"]}


class App:
    def __init__(self, directory, admin_key, secure_cookie=False, require_login=True):
        self.store = Store(directory)
        self.admin_key = admin_key
        self.secure_cookie = secure_cookie
        self.require_login = require_login
        self.lock = threading.Lock()
        self.sessions, self.uploads, self.previews, self.attempts = {}, {}, {}, {}

    def cache(self, collection, value):
        token = secrets.token_urlsafe(24)
        with self.lock:
            self.prune()
            if len(collection) >= 20:
                raise Error("待处理导入过多，请稍后再试", 429)
            collection[token] = (time.time() + 1800, value)
        return token

    def prune(self):
        for collection in (self.sessions, self.uploads, self.previews, self.attempts):
            for key, entry in list(collection.items()):
                if entry[0] < time.time():
                    del collection[key]

    def cached(self, collection, token):
        with self.lock:
            self.prune()
            entry = collection.get(token)
            if not entry:
                raise Error("导入预览已过期，请重新上传", 410)
            return entry[1]


def handler_for(app):
    class Handler(BaseHTTPRequestHandler):
        server_version = "Workbench"

        def log_message(self, *_args):
            pass  # Do not log credentials or project share URLs.

        def reply(self, value, status=200, cookie=None, content_type="application/json; charset=utf-8"):
            raw = value if isinstance(value, bytes) else dumps(value).encode()
            self.send_response(status)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(raw)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Referrer-Policy", "no-referrer")
            self.send_header("Content-Security-Policy", "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'none'; form-action 'self'")
            if cookie:
                self.send_header("Set-Cookie", cookie)
            self.end_headers()
            self.wfile.write(raw)

        def body(self):
            if "application/json" not in self.headers.get("Content-Type", ""):
                raise Error("请使用 JSON 请求", 415)
            try:
                size = int(self.headers.get("Content-Length", "0"))
            except ValueError:
                raise Error("无效请求长度")
            if size <= 0 or size > MAX_FILE * 4 // 3 + 8192:
                raise Error("请求体为空或过大", 413)
            try:
                body = json.loads(self.rfile.read(size))
                if not isinstance(body, dict):
                    raise ValueError()
                return body
            except (ValueError, UnicodeDecodeError) as exc:
                raise Error("无效 JSON 请求") from exc

        def local_request(self):
            host = urlparse("http://" + self.headers.get("Host", "")).hostname or ""
            try:
                local_host = host == "localhost" or ipaddress.ip_address(host).is_loopback
                local_client = ipaddress.ip_address(self.client_address[0]).is_loopback
            except ValueError:
                local_host = local_client = False
            if not local_host or not local_client:
                raise Error("当前为本机免登录模式，只允许本机地址访问", 403)

        def authorized(self):
            if not app.require_login:
                self.local_request()
                return "local"
            cookie = SimpleCookie()
            try:
                cookie.load(self.headers.get("Cookie", ""))
                token = cookie["wb_session"].value if "wb_session" in cookie else ""
            except Exception:
                token = ""
            with app.lock:
                app.prune()
                if token not in app.sessions:
                    raise Error("请先登录", 401)
            return token

        def do_GET(self):
            self.handle_request("GET")

        def do_POST(self):
            self.handle_request("POST")

        def handle_request(self, method):
            try:
                self.route(method)
            except Error as exc:
                self.reply({"error": str(exc)}, exc.status)
            except Exception:
                import traceback
                traceback.print_exc()
                self.reply({"error": "服务处理失败，请检查服务日志"}, 500)

        def route(self, method):
            if not app.require_login:
                self.local_request()
            path = urlparse(self.path).path
            parts = path.strip("/").split("/")
            if method == "GET" and path == "/healthz":
                return self.reply({"ok": True, "service": "project-workbench", "version": "0.3", "features": ["project-json"], "requireLogin": app.require_login})
            if method == "GET" and (path in ("/", "/app.js", "/style.css", "/report.js") or (len(parts) == 2 and parts[0] == "share")):
                asset = "index.html" if path == "/" or parts[0] == "share" else parts[0]
                mime = {"index.html": "text/html", "app.js": "text/javascript", "report.js": "text/javascript", "style.css": "text/css"}[asset]
                return self.reply((ROOT / "static" / asset).read_bytes(), content_type=mime + "; charset=utf-8")
            if method == "GET" and path in ("/report.schema.json", "/project.schema.json"):
                schema_path = ROOT.parent / "skills" / "ipd-report-json" / "references" / path.lstrip("/")
                return self.reply(schema_path.read_bytes())
            if method == "GET" and len(parts) == 3 and parts[:2] == ["api", "share"]:
                with app.store.connect() as db:
                    row = db.execute("SELECT id FROM projects WHERE share_hash=?", (digest(parts[2]),)).fetchone()
                if not row:
                    raise Error("分享链接不存在或已撤销", 404)
                snapshot = app.store.snapshot(row[0])
                snapshot["project"].pop("mapping", None)
                snapshot["history"] = []
                return self.reply(snapshot)
            if method == "POST":
                # Only same-origin JSON requests can mutate data. No permissive CORS.
                origin = self.headers.get("Origin")
                if origin and urlparse(origin).netloc != self.headers.get("Host"):
                    raise Error("不允许跨站请求", 403)
                body = self.body()
                if path == "/api/login":
                    ip = self.client_address[0]
                    with app.lock:
                        app.prune()
                        expiry, attempts = app.attempts.get(ip, (time.time() + 300, 0))
                        if attempts >= 10:
                            raise Error("尝试过多，请 5 分钟后重试", 429)
                        if not hmac.compare_digest(digest(text(body.get("key"))), digest(app.admin_key)):
                            app.attempts[ip] = (expiry, attempts + 1)
                            raise Error("管理口令不正确", 401)
                        app.attempts.pop(ip, None)
                        token = secrets.token_urlsafe(32)
                        app.sessions[token] = (time.time() + 8 * 3600, True)
                    cookie = f"wb_session={token}; HttpOnly; SameSite=Strict; Path=/; Max-Age=28800" + ("; Secure" if app.secure_cookie else "")
                    return self.reply({"ok": True}, cookie=cookie)
            session = self.authorized()
            if method == "GET" and path == "/api/session":
                return self.reply({"ok": True, "fields": FIELDS, "defaultTags": DEFAULT_TAGS, "requireLogin": app.require_login})
            if method == "POST" and path == "/api/reports/preview":
                preview = app.store.preview_report(body.get("report"))
                token = app.cache(app.previews, preview)
                return self.reply({k: v for k, v in preview.items() if k != "report"} | {"token": token})
            if method == "POST" and path == "/api/reports/commit":
                preview = app.cached(app.previews, body.get("token"))
                if preview.get("kind") != "report":
                    raise Error("请先预览项目 JSON")
                result = app.store.save_report(preview)
                with app.lock:
                    app.previews.pop(body.get("token"), None)
                return self.reply(result)
            if method == "POST" and path == "/api/logout":
                with app.lock:
                    app.sessions.pop(session, None)
                return self.reply({"ok": True}, cookie="wb_session=; HttpOnly; SameSite=Strict; Path=/; Max-Age=0")
            if path == "/api/projects":
                if method == "POST":
                    return self.reply(app.store.create(body), 201)
                with app.store.connect() as db:
                    projects = []
                    for row in db.execute("SELECT p.id,COUNT(i.id) AS count FROM projects p LEFT JOIN items i ON p.id=i.project_id GROUP BY p.id ORDER BY p.updated DESC"):
                        projects.append({**app.store.project(row["id"], db), "count": row["count"]})
                return self.reply(projects)
            if len(parts) < 3 or parts[:2] != ["api", "projects"]:
                raise Error("接口不存在", 404)
            project_id = parts[2]
            project = app.store.project(project_id)
            action = parts[3] if len(parts) == 4 else ""
            if method == "POST" and action == "problems":
                return self.reply(app.store.edit_problem(project_id, body))
            if method == "POST" and action == "delete":
                result = app.store.delete(project_id, body)
                with app.lock:
                    for collection in (app.uploads, app.previews):
                        for token, (_, value) in list(collection.items()):
                            if value.get("projectId") == project_id:
                                collection.pop(token, None)
                return self.reply(result)
            if not action and len(parts) == 3:
                return self.reply(app.store.snapshot(project_id) if method == "GET" else app.store.update(project_id, body))
            if method == "GET" and action == "report":
                snapshot = app.store.snapshot(project_id)
                if "report" not in snapshot:
                    raise Error("该项目不是 JSON 拆解项目")
                return self.reply(snapshot["report"])
            if method == "POST" and action in ("upload", "preview", "commit") and project.get("mode") == "report":
                raise Error("此项目请导入完整 JSON 更新")
            if method == "POST" and action == "upload":
                try:
                    raw = base64.b64decode(body.get("content", ""), validate=True)
                except (ValueError, TypeError) as exc:
                    raise Error("文件编码无效") from exc
                filename = text(body.get("filename"))[:200]
                headers, records, sheets, sheet = parse_table(filename, raw, body.get("sheet"))
                mapping = {}
                for field, aliases in ALIASES.items():
                    saved = project["mapping"].get(field)
                    mapping[field] = saved if saved in headers else next((h for h in headers if h.lower() in [a.lower() for a in aliases]), "")
                token = app.cache(app.uploads, {"projectId": project_id, "filename": filename, "records": records})
                return self.reply({"token": token, "headers": headers, "sample": [r for _, r in records[:5]], "total": len(records), "sheets": sheets, "sheet": sheet, "mapping": mapping})
            if method == "POST" and action == "preview":
                upload = app.cached(app.uploads, body.get("token"))
                if upload["projectId"] != project_id:
                    raise Error("导入项目不匹配", 403)
                preview = app.store.preview(project_id, upload["filename"], upload["records"], body.get("mapping"))
                token = app.cache(app.previews, preview)
                return self.reply({k: v for k, v in preview.items() if k != "items"} | {"token": token})
            if method == "POST" and action == "commit":
                preview = app.cached(app.previews, body.get("token"))
                if preview.get("kind") == "report" or preview["projectId"] != project_id:
                    raise Error("导入项目不匹配", 403)
                result = app.store.commit(preview)
                with app.lock:
                    app.previews.pop(body.get("token"), None)
                return self.reply(result)
            if method == "POST" and action == "share":
                token = secrets.token_urlsafe(32) if body.get("enabled") else None
                with app.store.connect() as db:
                    db.execute("UPDATE projects SET share_hash=? WHERE id=?", (digest(token) if token else None, project_id))
                return self.reply({"path": "/share/" + token if token else None, "feed": "/api/share/" + token if token else None})
            raise Error("接口不存在", 404)

    return Handler


def main():
    directory = Path(os.environ.get("WORKBENCH_DATA", str(ROOT / "data")))
    directory.mkdir(parents=True, exist_ok=True)
    key = os.environ.get("WORKBENCH_ADMIN_KEY", "")
    if not key:
        keyfile = directory / "admin.key"
        if not keyfile.exists():
            keyfile.write_text(secrets.token_urlsafe(24), encoding="utf-8")
            try:
                keyfile.chmod(0o600)
            except OSError:
                pass
        key = keyfile.read_text(encoding="utf-8").strip()
    if len(key) < 16:
        raise SystemExit("管理口令至少 16 个字符")
    require_login = os.environ.get("WORKBENCH_REQUIRE_LOGIN", "1") != "0"
    app = App(directory, key, os.environ.get("WORKBENCH_SECURE_COOKIE") == "1", require_login=require_login)
    host, port = os.environ.get("WORKBENCH_HOST", "127.0.0.1"), int(os.environ.get("WORKBENCH_PORT", "8765"))
    if not require_login and host not in ("127.0.0.1", "localhost"):
        raise SystemExit("免登录模式只允许绑定 127.0.0.1 或 localhost")
    server = ThreadingHTTPServer((host, port), handler_for(app))
    server.daemon_threads = True
    print(f"Workbench ready: http://{host}:{port}", flush=True)
    print(f"Admin key: {'environment WORKBENCH_ADMIN_KEY' if os.environ.get('WORKBENCH_ADMIN_KEY') else directory / 'admin.key'}", flush=True)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
